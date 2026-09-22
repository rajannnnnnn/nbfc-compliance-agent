"""bytes/text -> ParsedDocument (pdf, txt, docx, ocr fallback). LLD §3.1 (module map) / HLD §3.1.

No branch writes bytes to disk or returns a file path — CLAUDE.md §2.3: no document bytes
persisted, not on disk, not anywhere. This module returns text in memory only.
"""

import io

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.config import Settings
from app.domain.documents import ParsedDocument


class ParseError(Exception):
    pass


def _parse_pdf_text_layer(data: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages), len(pages)


def _parse_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _ocr_fallback(data: bytes, page_count: int) -> str:
    """OCR is not wired to a real engine in this build (no scanned-document fixtures exist
    yet to validate against) — this raises rather than silently returning empty text, which
    would look like a successful extraction of nothing. Wire pytesseract or a hosted OCR API
    here when a real scanned-document fixture is available."""
    raise ParseError(
        f"document requires OCR ({page_count} pages, low character density) — "
        "no OCR engine is wired in this build"
    )


def parse_document(
    *, filename: str, data: bytes | None = None, text: str | None = None, settings: Settings
) -> ParsedDocument:
    """Exactly one of `data` (PDF/DOCX bytes) or `text` (already-extracted plain text, the
    common API path per LLD §15.1 where the submission carries `text` directly) must be
    given."""
    if text is not None:
        return ParsedDocument(text=text, char_count=len(text), page_count=None, ocr_used=False)

    if data is None:
        raise ParseError("either data or text must be supplied")

    lower = filename.lower()
    if lower.endswith(".pdf"):
        extracted, page_count = _parse_pdf_text_layer(data)
        chars_per_page = len(extracted) / max(page_count, 1)
        if chars_per_page < settings.ocr_char_per_page_threshold:
            extracted = _ocr_fallback(data, page_count)
            return ParsedDocument(
                text=extracted, char_count=len(extracted), page_count=page_count, ocr_used=True
            )
        return ParsedDocument(
            text=extracted, char_count=len(extracted), page_count=page_count, ocr_used=False
        )

    if lower.endswith(".docx"):
        extracted = _parse_docx(data)
        return ParsedDocument(
            text=extracted, char_count=len(extracted), page_count=None, ocr_used=False
        )

    if lower.endswith(".txt"):
        extracted = data.decode("utf-8", errors="replace")
        return ParsedDocument(
            text=extracted, char_count=len(extracted), page_count=None, ocr_used=False
        )

    raise ParseError(f"unsupported file type: {filename}")
